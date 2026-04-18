"""Unit tests for pdfkit.converter.PDFConverter."""

from __future__ import annotations

from pathlib import Path

import pytest

from pdfkit import SUPPORTED_FORMATS, PDFConverter


def test_rejects_missing_file(tmp_path):
    with pytest.raises(FileNotFoundError):
        PDFConverter(tmp_path / "nope.pdf")


def test_rejects_non_pdf_extension(tmp_path):
    fake = tmp_path / "not_a.pdf.txt"
    fake.write_text("hi")
    with pytest.raises(ValueError):
        PDFConverter(fake)


def test_rejects_unknown_format(sample_pdf, tmp_path):
    with pytest.raises(ValueError):
        PDFConverter(sample_pdf).convert(tmp_path, formats=["rtf"])


def test_txt_extracts_text(sample_pdf, tmp_path):
    out = PDFConverter(sample_pdf).convert(tmp_path, formats=["txt"])["txt"]
    text = Path(out).read_text(encoding="utf-8")
    assert "PDFKit Sample Document" in text
    assert "Second Page" in text
    assert "Alice" in text and "95" in text


def test_html_preserves_styling_hints(sample_pdf, tmp_path):
    out = PDFConverter(sample_pdf).convert(tmp_path, formats=["html"])["html"]
    html = Path(out).read_text(encoding="utf-8")
    # PyMuPDF emits inline CSS for font family and size on every span.
    assert "<html" in html and "</html>" in html
    assert "font-family" in html
    assert "font-size" in html
    assert "PDFKit Sample Document" in html


def test_markdown_contains_body(sample_pdf, tmp_path):
    out = PDFConverter(sample_pdf).convert(tmp_path, formats=["md"])["md"]
    md = Path(out).read_text(encoding="utf-8")
    assert "PDFKit Sample Document" in md
    assert "Second Page" in md
    # Page separator from our converter.
    assert "page 1" in md and "page 2" in md


def test_docx_is_valid_zip(sample_pdf, tmp_path):
    import zipfile

    out = PDFConverter(sample_pdf).convert(tmp_path, formats=["docx"])["docx"]
    assert zipfile.is_zipfile(out)
    with zipfile.ZipFile(out) as zf:
        names = zf.namelist()
        assert "word/document.xml" in names
        body = zf.read("word/document.xml").decode("utf-8", errors="ignore")
        assert "PDFKit Sample Document" in body


def test_xlsx_has_table_sheet(sample_pdf, tmp_path):
    from openpyxl import load_workbook

    out = PDFConverter(sample_pdf).convert(tmp_path, formats=["xlsx"])["xlsx"]
    wb = load_workbook(out)
    assert len(wb.sheetnames) >= 1
    found = False
    for name in wb.sheetnames:
        rows = [
            [c.value for c in row]
            for row in wb[name].iter_rows()
        ]
        flat = [str(v) for row in rows for v in row if v is not None]
        if "Alice" in flat and "95" in flat:
            found = True
            break
    assert found, f"table not found; sheets={wb.sheetnames}"


def test_page_range_limits_output(sample_pdf, tmp_path):
    out = PDFConverter(sample_pdf).convert(
        tmp_path, formats=["txt"], pages=(0, 1)
    )["txt"]
    text = Path(out).read_text(encoding="utf-8")
    assert "PDFKit Sample Document" in text
    assert "Second Page" not in text


def test_empty_page_range_raises(sample_pdf, tmp_path):
    with pytest.raises(ValueError):
        PDFConverter(sample_pdf).convert(
            tmp_path, formats=["txt"], pages=(2, 2)
        )


def test_convert_all_formats(sample_pdf, tmp_path):
    outputs = PDFConverter(sample_pdf).convert(
        tmp_path, formats=list(SUPPORTED_FORMATS)
    )
    assert set(outputs) == set(SUPPORTED_FORMATS)
    for path in outputs.values():
        assert Path(path).exists()
        assert Path(path).stat().st_size > 0
